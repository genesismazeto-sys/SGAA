from __future__ import annotations

import csv
import datetime as dt
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Iterable

import openpyxl

from app.text import normalize_header


SUPPORTED_STUDENT_IMPORT_EXTENSIONS = {"csv", "xlsx", "xls"}
CANONICAL_STUDENT_IMPORT_HEADERS = ("Aluno", "E-mail", "Matricula")
_HEADER_ALIASES = (
    {"aluno", "nome"},
    {"e-mail", "email"},
    {"matricula"},
)


class StudentImportError(ValueError):
    pass


@dataclass(frozen=True)
class StudentImportRow:
    aluno: str
    email: str
    matricula: str
    source_row: int


@dataclass(frozen=True)
class _Cell:
    value: object
    number_format: str = ""
    kind: str = "text"


def _numeric_text(value: int | float | Decimal, number_format: str = "") -> str:
    decimal_value = Decimal(str(value))
    if not decimal_value.is_finite():
        raise StudentImportError("A matrícula contém um valor numérico inválido.")
    significant_digits = len(decimal_value.normalize().as_tuple().digits)
    if significant_digits > 15:
        raise StudentImportError(
            "A matrícula contém um valor numérico sem representação textual segura."
        )
    primary_format = str(number_format or "").split(";", 1)[0].strip()
    if decimal_value == decimal_value.to_integral_value() and re.fullmatch(r"0+", primary_format):
        return str(int(decimal_value)).zfill(len(primary_format))

    rendered = format(decimal_value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered


def _cell_text(cell: _Cell, *, identifier: bool = False) -> str:
    value = cell.value
    if value is None:
        return ""
    if identifier and (
        cell.kind in {"date", "datetime"}
        or isinstance(value, (dt.date, dt.datetime, dt.time))
    ):
        raise StudentImportError("A matrícula não pode ser uma data ou hora.")
    if identifier and cell.kind in {"error", "formula"}:
        raise StudentImportError("A matrícula não pode conter fórmula ou valor de erro.")
    if identifier and cell.kind == "boolean":
        raise StudentImportError("A matrícula não pode ser um valor booleano.")
    if identifier and isinstance(value, (bytes, bytearray, list, tuple, dict, set)):
        raise StudentImportError("A matrícula deve ser um identificador escalar.")
    if isinstance(value, bool):
        return "VERDADEIRO" if value else "FALSO"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return _numeric_text(value, cell.number_format if identifier else "")
    return str(value).strip()


def _row_has_content(cells: list[_Cell]) -> bool:
    return any(_cell_text(cell) for cell in cells)


def _validate_header(cells: list[_Cell], row_number: int) -> None:
    if len(cells) != 3:
        raise StudentImportError(
            f"Linha {row_number}: o arquivo deve ter exatamente as colunas Aluno, E-mail, Matricula."
        )
    header = tuple(normalize_header(_cell_text(cell)) for cell in cells[:3])
    if len(header) != 3 or any(value not in aliases for value, aliases in zip(header, _HEADER_ALIASES)):
        raise StudentImportError(
            "O cabeçalho deve seguir exatamente esta ordem: Aluno, E-mail, Matricula."
        )


def _normalize_data_rows(
    rows: Iterable[tuple[int, list[_Cell]]],
    *,
    allow_empty: bool = False,
) -> list[StudentImportRow]:
    normalized: list[StudentImportRow] = []
    seen_emails: dict[str, int] = {}
    seen_matriculas: dict[str, int] = {}

    for row_number, cells in rows:
        if not _row_has_content(cells):
            continue
        if len(cells) > 3:
            raise StudentImportError(
                f"Linha {row_number}: colunas extras não são permitidas; use Aluno, E-mail, Matricula."
            )

        aluno = _cell_text(cells[0]) if len(cells) > 0 else ""
        email = _cell_text(cells[1]) if len(cells) > 1 else ""
        matricula = _cell_text(cells[2], identifier=True) if len(cells) > 2 else ""
        missing = [
            label
            for label, value in (("Aluno", aluno), ("E-mail", email), ("Matricula", matricula))
            if not value
        ]
        if missing:
            raise StudentImportError(f"Linha {row_number}: campo(s) obrigatório(s) ausente(s): {', '.join(missing)}.")

        email_key = email.casefold()
        if email_key in seen_emails:
            raise StudentImportError(
                f"Linha {row_number}: e-mail duplicado no arquivo; primeira ocorrência na linha {seen_emails[email_key]}."
            )
        if matricula in seen_matriculas:
            raise StudentImportError(
                f"Linha {row_number}: matrícula duplicada no arquivo; primeira ocorrência na linha {seen_matriculas[matricula]}."
            )
        seen_emails[email_key] = row_number
        seen_matriculas[matricula] = row_number
        normalized.append(StudentImportRow(aluno, email, matricula, row_number))

    if not normalized and not allow_empty:
        raise StudentImportError("O arquivo não contém nenhum aluno para importar.")
    return normalized


def _normalize_rows(rows: Iterable[tuple[int, list[_Cell]]]) -> list[StudentImportRow]:
    iterator = iter(rows)
    for row_number, cells in iterator:
        if not _row_has_content(cells):
            continue
        _validate_header(cells, row_number)
        return _normalize_data_rows(iterator)
    raise StudentImportError("O arquivo está vazio.")


def _csv_rows(path: Path) -> Iterable[tuple[int, list[_Cell]]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            sample = source.read(4096)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            for row_number, values in enumerate(
                csv.reader(source, dialect=dialect, strict=True),
                start=1,
            ):
                yield row_number, [_Cell(value) for value in values]
    except UnicodeDecodeError as exc:
        raise StudentImportError("O CSV deve estar codificado em UTF-8.") from exc
    except csv.Error as exc:
        raise StudentImportError("O CSV está malformado e não pôde ser lido.") from exc


def _xlsx_rows(path: Path) -> Iterable[tuple[int, list[_Cell]]]:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise StudentImportError("A planilha XLSX está malformada e não pôde ser lida.") from exc
    rows = []
    try:
        sheet = workbook.active
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            kinds = {"d": "date", "e": "error", "f": "formula", "b": "boolean"}
            rows.append(
                (
                    row_number,
                    [
                        _Cell(cell.value, cell.number_format, kinds.get(cell.data_type, cell.data_type))
                        for cell in cells
                    ],
                )
            )
    finally:
        workbook.close()
    return rows


def _xls_rows(path: Path) -> Iterable[tuple[int, list[_Cell]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise StudentImportError("A leitura de arquivos XLS não está disponível.") from exc

    try:
        workbook = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    except Exception as exc:
        raise StudentImportError("A planilha XLS está malformada e não pôde ser lida.") from exc
    rows = []
    try:
        if workbook.nsheets == 0:
            return rows
        sheet = workbook.sheet_by_index(0)
        for row_index in range(sheet.nrows):
            cells: list[_Cell] = []
            for column_index in range(sheet.ncols):
                source_cell = sheet.cell(row_index, column_index)
                number_format = ""
                if source_cell.ctype == xlrd.XL_CELL_NUMBER:
                    xf = workbook.xf_list[source_cell.xf_index]
                    fmt = workbook.format_map.get(xf.format_key)
                    number_format = fmt.format_str if fmt is not None else ""
                kinds = {
                    xlrd.XL_CELL_DATE: "date",
                    xlrd.XL_CELL_ERROR: "error",
                    xlrd.XL_CELL_BOOLEAN: "boolean",
                    xlrd.XL_CELL_NUMBER: "number",
                    xlrd.XL_CELL_TEXT: "text",
                }
                cells.append(
                    _Cell(
                        source_cell.value,
                        number_format,
                        kinds.get(source_cell.ctype, "unknown"),
                    )
                )
            rows.append((row_index + 1, cells))
    finally:
        workbook.release_resources()
    return rows
def parse_student_import(path: str | Path) -> list[StudentImportRow]:
    source_path = Path(path)
    extension = source_path.suffix.lower().lstrip(".")
    if extension not in SUPPORTED_STUDENT_IMPORT_EXTENSIONS:
        raise StudentImportError("Formato não suportado. Envie um arquivo CSV, XLSX ou XLS.")
    if extension == "csv":
        rows = _csv_rows(source_path)
    elif extension == "xlsx":
        rows = _xlsx_rows(source_path)
    else:
        rows = _xls_rows(source_path)
    try:
        return _normalize_rows(rows)
    finally:
        close = getattr(rows, "close", None)
        if close:
            close()


def normalize_student_import_values(
    values: Iterable[tuple[object, object, object]],
    *,
    allow_empty: bool = False,
) -> list[StudentImportRow]:
    rows = (
        (row_number, [_Cell(aluno), _Cell(email), _Cell(matricula)])
        for row_number, (aluno, email, matricula) in enumerate(values, start=1)
    )
    return _normalize_data_rows(rows, allow_empty=allow_empty)


__all__ = [
    "CANONICAL_STUDENT_IMPORT_HEADERS",
    "SUPPORTED_STUDENT_IMPORT_EXTENSIONS",
    "StudentImportError",
    "StudentImportRow",
    "normalize_student_import_values",
    "parse_student_import",
]
