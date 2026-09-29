"""Spreadsheet formulas and error cells are never student data (UI-B34).

The student import is a data import, not a spreadsheet engine. A formula or an
error cell in any of the three student columns (Aluno, E-mail, Matricula)
rejects the file at the true source row, in the server parser and in the
browser preview alike, by the cell TYPE the reader reports -- never by its
display text, never by evaluating the formula and never by using the value
Excel cached for it. Before, the server imported an XLSX formula as its formula
text (``=UPPER("bia")`` became the name) and an XLS error cell as its code
(``#N/A`` became "42"), while the preview showed the cached value as valid.

CSV has no cell types: a literal ``=...`` or ``#N/A`` stays ordinary text.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pytest
import xlwt

import main
from app.password_tokens import PURPOSE_FIRST_ACCESS, issue_password_token
from app.student_import import StudentImportError, parse_student_import
from app.user_accounts import create_usuario_pending
from tests.canonical_matrix_test_support import login_admin
from tests.test_student_import_optional_header import _browser_preview
from tests.versioned_test_support import isolated_versioned_app_env


@dataclass(frozen=True)
class F:
    """A formula cell; ``cached`` is the value Excel would have saved with it."""

    formula: str
    cached: str | None = None


@dataclass(frozen=True)
class E:
    """An error cell (#N/A)."""


HEADER = ["Aluno", "E-mail", "Matricula"]
LABELS = HEADER
VALID = [
    ["Primeira Pessoa", "primeira@example.com", "001234"],
    ["Segunda Pessoa", "segunda@example.com", "ACE-0002"],
    ["Terceira Pessoa", "terceira@example.com", "T-0003"],
]
FORMULAS = (
    F('UPPER("bia")', "BIA"),
    F('"bia"&"@example.com"', "bia@example.com"),
    F('"00"&"1234"', "001234"),
)


def _message(row: int, column: int) -> str:
    return (
        f"Linha {row}: fórmula ou valor de erro na coluna {LABELS[column]}; "
        "a importação aceita apenas valores digitados."
    )


def _xlsx_bytes(rows) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    cached: dict[str, str] = {}
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(value, F):
                cell.value = f"={value.formula}"
                if value.cached is not None:
                    cached[cell.coordinate] = value.cached
            elif isinstance(value, E):
                cell.value = "#N/A"  # openpyxl stores it as an error cell
                assert cell.data_type == "e"
            else:
                cell.value = value
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    if not cached:
        return payload.getvalue()

    # openpyxl never writes a formula's cached result; Excel always does. Put
    # it back so the file looks exactly like one saved by Excel.
    source = zipfile.ZipFile(io.BytesIO(payload.getvalue()))
    rebuilt = io.BytesIO()
    with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for coordinate, value in cached.items():
                    xml, count = re.subn(
                        rf'<c r="{coordinate}"><f>(.*?)</f><v\s*/></c>',
                        lambda match: (
                            f'<c r="{coordinate}" t="str"><f>{match.group(1)}</f>'
                            f"<v>{value}</v></c>"
                        ),
                        xml,
                    )
                    assert count == 1, coordinate
                data = xml.encode("utf-8")
            target.writestr(item, data)
    return rebuilt.getvalue()


def _xls_bytes(rows) -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Alunos")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if isinstance(value, F):
                sheet.write(row_index, column_index, xlwt.Formula(value.formula))
            elif isinstance(value, E):
                sheet.row(row_index).set_cell_error(column_index, 0x2A)  # #N/A
            else:
                sheet.write(row_index, column_index, value)
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _csv_bytes(rows) -> bytes:
    text = io.StringIO()
    csv.writer(text, lineterminator="\n").writerows(rows)
    return text.getvalue().encode("utf-8")


BUILDERS = {"xlsx": _xlsx_bytes, "xls": _xls_bytes, "csv": _csv_bytes}


def _write(tmp_path: Path, extension: str, rows, name: str) -> Path:
    path = tmp_path / f"{name}.{extension}"
    path.write_bytes(BUILDERS[extension](rows))
    return path


def _server(path: Path) -> dict:
    try:
        rows = parse_student_import(path)
    except StudentImportError as exc:
        return {"ok": False, "error": str(exc)}
    return {
        "ok": True,
        "rows": [{"aluno": r.aluno, "email": r.email, "matricula": r.matricula} for r in rows],
    }


def _ok(rows) -> dict:
    return {"ok": True, "rows": [{"aluno": a, "email": e, "matricula": m} for a, e, m in rows]}


def _assert_both(paths, expected):
    assert [_server(path) for path in paths] == expected
    assert _browser_preview(paths) == expected


def _with(row_values, column, value):
    row = list(row_values)
    row[column] = value
    return row


# --- server and preview reject by cell type, at the true row -----------------


@pytest.mark.parametrize("extension", ("xlsx", "xls"))
@pytest.mark.parametrize("with_header", (True, False), ids=("with-header", "without-header"))
def test_formula_in_any_student_column_rejects_at_the_true_row(tmp_path, extension, with_header):
    top = [HEADER] if with_header else []
    paths, expected = [], []
    for column in range(3):
        for position in range(3):  # first, middle, last data row
            rows = [list(row) for row in VALID]
            rows[position][column] = FORMULAS[column]
            paths.append(_write(tmp_path, extension, [*top, *rows], f"f{column}{position}"))
            expected.append({"ok": False, "error": _message(len(top) + position + 1, column)})
    _assert_both(paths, expected)


@pytest.mark.parametrize("extension", ("xlsx", "xls"))
def test_error_cell_in_any_student_column_rejects_at_the_true_row(tmp_path, extension):
    paths, expected = [], []
    for column in range(3):
        rows = [HEADER, VALID[0], _with(VALID[1], column, E()), VALID[2]]
        paths.append(_write(tmp_path, extension, rows, f"e{column}"))
        expected.append({"ok": False, "error": _message(3, column)})
    _assert_both(paths, expected)


def test_xlsx_cached_values_that_look_valid_are_never_used(tmp_path):
    """The preview used to show the cached "BIA" / "bia@example.com" /
    "001234" as a valid student while the server imported the formula text."""
    paths, expected = [], []
    for column in range(3):
        rows = [HEADER, _with(VALID[0], column, FORMULAS[column]), VALID[1]]
        paths.append(_write(tmp_path, "xlsx", rows, f"cached{column}"))
        expected.append({"ok": False, "error": _message(2, column)})
    _assert_both(paths, expected)


def test_the_same_values_typed_as_plain_cells_still_import(tmp_path):
    typed = [["BIA", "bia@example.com", "001234"], VALID[1]]
    paths = [
        _write(tmp_path, extension, rows, f"typed-{extension}-{len(rows)}")
        for extension in ("xlsx", "xls", "csv")
        for rows in ([HEADER, *typed], typed)
    ]
    _assert_both(paths, [_ok(typed)] * len(paths))


@pytest.mark.parametrize("extension", ("xlsx", "xls"))
def test_formula_header_is_never_a_header_and_blank_result_is_still_content(tmp_path, extension):
    formula_header = [F('"Aluno"', "Aluno"), F('"E-mail"', "E-mail"), F('"Matricula"', "Matricula")]
    blank_result = [F('""', ""), "", ""]
    paths = [
        _write(tmp_path, extension, [formula_header, *VALID], "formula-header"),
        _write(tmp_path, extension, [HEADER, VALID[0], blank_result, VALID[1]], "blank-result"),
    ]
    _assert_both(
        paths,
        [
            {"ok": False, "error": _message(1, 0)},
            {"ok": False, "error": _message(3, 0)},
        ],
    )


def test_csv_literal_text_keeps_its_existing_meaning(tmp_path):
    """CSV has no formula cells: "=..." and "#N/A" are plain text, as before."""
    literal = [
        ['=UPPER("bia")', "bia@example.com", "B-0002"],
        ["#N/A", "na@example.com", "=1+1"],
        ["Cid", '="cid"&"@example.com"', "001234"],
    ]
    path = _write(tmp_path, "csv", [HEADER, *literal], "literal")
    _assert_both([path], [_ok(literal)])


# --- nothing is written: modal import, Editar Turma and Nova Turma ------------

TURMA_ID = 1  # seeded PPA-T10
TABLES = ("alunos", "usuarios", "usuario_credenciais", "senha_tokens", "turmas")


def _seed_existing_student() -> None:
    with main.app.app_context():
        conn = main.get_db_connection()
        usuario_id = create_usuario_pending(
            conn, "Aluna Existente", "existente@example.com", "aluno"
        ).lastrowid
        conn.execute(
            "INSERT INTO alunos (usuario_id,nome,email,matricula,turma_id,status)"
            " VALUES (?,?,?,?,?,'Ativo')",
            (usuario_id, "Aluna Existente", "existente@example.com", "EX-0001", TURMA_ID),
        )
        issue_password_token(conn, usuario_id, PURPOSE_FIRST_ACCESS)
        conn.commit()


def _database_state():
    with main.app.app_context():
        conn = main.get_db_connection()
        return {
            table: [tuple(row) for row in conn.execute(f"SELECT * FROM {table} ORDER BY 1")]
            for table in TABLES
        }


def _import_rows(position: int):
    # Row 1 would update the existing student's e-mail, the others would create
    # students; the formula sits at the given data position.
    rows = [
        ["Aluna Existente", "existente.nova@example.com", "EX-0001"],
        ["Nova Pessoa", "nova@example.com", "NEW-0001"],
        ["Outra Pessoa", "outra@example.com", "NEW-0002"],
    ]
    rows[position][0] = FORMULAS[0]
    return [HEADER, *rows], _message(position + 2, 0)


def _uploaded_files(tmp_path: Path):
    return sorted(path for path in (tmp_path / "uploads").rglob("*") if path.is_file())


def _turma_form(curso_id: int, numero: int) -> dict:
    return {
        "curso_id": str(curso_id),
        "numero_turma": str(numero),
        "ano_inicio": "2025",
        "semestre_inicio": "2",
        "turno": "Noite",  # would change if anything were committed
        "status": "Ativa",
        "matriz_id": "",
        "aluno_nome[]": "Aluna Existente",
        "aluno_email[]": "existente@example.com",
        "aluno_matricula[]": "EX-0001",
        "aluno_situacao[]": "ATIVO",
        "aluno_importado[]": "0",
    }


@pytest.mark.parametrize("extension", ("xlsx", "xls"))
@pytest.mark.parametrize("position", (0, 1, 2), ids=("first", "middle", "last"))
@pytest.mark.parametrize("route", ("modal", "editar", "nova"))
def test_formula_file_writes_nothing_and_leaves_no_upload(tmp_path, extension, position, route):
    with isolated_versioned_app_env(tmp_path, f"expr-{route}.db") as env:
        _seed_existing_student()
        login_admin(env["client"])
        rows, error = _import_rows(position)
        upload = (io.BytesIO(BUILDERS[extension](rows)), f"alunos.{extension}")
        with main.app.app_context():
            curso_id = main.get_db_connection().execute(
                "SELECT curso_id FROM turmas WHERE id=?", (TURMA_ID,)
            ).fetchone()["curso_id"]
        before = _database_state()

        if route == "modal":
            response = env["client"].post(
                "/admin/turmas/importar",
                data={"turma_id": str(TURMA_ID), "csv_arquivo": upload},
            )
            assert response.status_code == 302
            with env["client"].session_transaction() as session:
                flashes = session.get("_flashes", [])
            assert any(category == "error" and error in message for category, message in flashes)
        else:
            form = _turma_form(curso_id, 10 if route == "editar" else 77)
            form["student_import_file"] = upload
            url = f"/admin/editar_turma/{TURMA_ID}" if route == "editar" else "/admin/adicionar_turma"
            response = env["client"].post(url, data=form)
            assert response.status_code == 200
            prefix = "Erro ao atualizar turma" if route == "editar" else "Erro ao criar turma"
            assert f"{prefix}: {error}" in response.get_data(as_text=True)

        assert _database_state() == before
        assert _uploaded_files(tmp_path) == []
