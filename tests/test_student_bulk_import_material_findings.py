from __future__ import annotations

import datetime as dt
import io
import uuid
import zipfile
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import xlwt
from flask import session
from werkzeug.datastructures import MultiDict

from app import db as app_db
from app.student_import import (
    StudentImportError,
    normalize_student_import_values,
    parse_student_import,
)
from app.views.admin import alunos_turmas_cursos as target
import main


@pytest.fixture(autouse=True)
def _isolated_student_import_database(tmp_path, monkeypatch):
    database_path = tmp_path / "student-import-material.db"
    monkeypatch.setattr(app_db, "DATABASE", str(database_path))
    monkeypatch.setitem(main.app.config, "DATABASE_PATH", str(database_path))
    with main.app.app_context():
        app_db.init_db()
    yield
    with main.app.app_context():
        app_db.close_db_connection(None)


def _seed_school(*, with_turma: bool = True) -> tuple[int, int, int | None]:
    suffix = uuid.uuid4().hex[:8].upper()
    with main.app.app_context():
        conn = app_db.get_db_connection()
        curso_id = conn.execute(
            "INSERT INTO cursos (nome,codigo,duracao_periodos,status) VALUES (?,?,8,'ativo') RETURNING id",
            (f"Curso material {suffix}", f"MF{suffix}"),
        ).fetchone()["id"]
        matriz_id = conn.execute(
            "INSERT INTO matrizes_atividades (curso_id,nome,status) VALUES (?,?,'ativa') RETURNING id",
            (curso_id, f"Matriz material {suffix}"),
        ).fetchone()["id"]
        turma_id = None
        if with_turma:
            turma_id = conn.execute(
                """
                INSERT INTO turmas
                    (nome,codigo,numero,curso_id,matriz_id,status,ano_inicio,semestre_inicio)
                VALUES (?,?,1,?,?,'Ativa',2026,1) RETURNING id
                """,
                (f"MF{suffix}-T1", f"MF{suffix}-T1", curso_id, matriz_id),
            ).fetchone()["id"]
        conn.commit()
        return int(curso_id), int(matriz_id), int(turma_id) if turma_id else None


def _insert_identity(
    *,
    name: str,
    email: str,
    user_type: str,
    matricula: str | None = None,
    turma_id: int | None = None,
    status: str = "Ativo",
) -> tuple[int, int | None]:
    with main.app.app_context():
        conn = app_db.get_db_connection()
        user_id = conn.execute(
            """
            INSERT INTO usuarios (nome,email,senha,tipo,nivel_acesso)
            VALUES (?,?,?,?,'administrativo') RETURNING id
            """,
            (name, email, "hash", user_type),
        ).fetchone()["id"]
        aluno_id = None
        if matricula is not None:
            aluno_id = conn.execute(
                """
                INSERT INTO alunos (usuario_id,nome,email,matricula,turma_id,status)
                VALUES (?,?,?,?,?,?) RETURNING id
                """,
                (user_id, name, email, matricula, turma_id, status),
            ).fetchone()["id"]
        conn.commit()
        return int(user_id), int(aluno_id) if aluno_id else None


def _post_modal(
    turma_id: int,
    rows: list[tuple[str, str, str]],
    *,
    on_conflict: str = "update",
):
    body = "Aluno,E-mail,Matricula\n" + "".join(
        f"{name},{email},{matricula}\n" for name, email, matricula in rows
    )
    with main.app.test_request_context(
        "/admin/turmas/importar",
        method="POST",
        data={
            "turma_id": str(turma_id),
            "csv_arquivo": (io.BytesIO(body.encode("utf-8")), "material.csv"),
            "on_conflict": on_conflict,
        },
    ):
        session["user_id"] = 999_999
        session["user_type"] = "admin"
        response = target.admin_turmas_importar()
        assert not list(
            (Path(main.app.config["UPLOAD_FOLDER"]) / "turmas_imports").glob("*")
        )
        return response, list(session.get("_flashes", []))


def _post_turma_form(
    mode: str,
    curso_id: int,
    matriz_id: int,
    turma_id: int | None,
    rows: list[tuple[str, str, str, str]],
    *,
    import_file: tuple[bytes, str] | None = None,
    imported_flags: list[str] | None = None,
):
    numero = 1 if turma_id else 2
    items = [
        ("curso_id", str(curso_id)),
        ("matriz_id", str(matriz_id)),
        ("numero_turma", str(numero)),
        ("ano_inicio", "2026"),
        ("semestre_inicio", "1"),
        ("turno", "Noite"),
        ("status", "Ativa"),
    ]
    if imported_flags is None:
        imported_flags = ["1"] * len(rows)
    assert len(imported_flags) == len(rows)
    for (name, email, matricula, status), imported_flag in zip(rows, imported_flags):
        items.extend(
            [
                ("aluno_nome[]", name),
                ("aluno_email[]", email),
                ("aluno_matricula[]", matricula),
                ("aluno_situacao[]", status),
                ("aluno_importado[]", imported_flag),
            ]
        )
    if import_file is None:
        body = "Aluno,E-mail,Matricula\n" + "".join(
            f"{name},{email},{matricula}\n" for name, email, matricula, _ in rows
        )
        import_file = (body.encode("utf-8"), "material-form.csv")
    items.append(
        ("student_import_file", (io.BytesIO(import_file[0]), import_file[1]))
    )
    path = "/admin/adicionar_turma" if mode == "add" else f"/admin/editar_turma/{turma_id}"
    with main.app.test_request_context(path, method="POST", data=MultiDict(items)):
        session["user_id"] = 999_999
        session["user_type"] = "admin"
        response = (
            target.admin_adicionar_turma()
            if mode == "add"
            else target.admin_editar_turma(int(turma_id))
        )
        assert not list(
            (Path(main.app.config["UPLOAD_FOLDER"]) / "turmas_imports").glob("*")
        )
        return response, list(session.get("_flashes", []))


def _xlsx_file(
    tmp_path: Path,
    value,
    *,
    data_type: str | None = None,
    number_format: str = "General",
    name: str = "Typed XLSX",
    email: str = "typed.xlsx@example.com",
) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Aluno", "E-mail", "Matricula"])
    sheet.append([name, email, value])
    cell = sheet.cell(2, 3)
    if data_type:
        cell.data_type = data_type
    if isinstance(value, (dt.date, dt.datetime, dt.time)) and number_format == "General":
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    else:
        cell.number_format = number_format
    path = tmp_path / f"typed-{uuid.uuid4().hex}.xlsx"
    workbook.save(path)
    return path


def _xlsx_float_artifact_file(tmp_path: Path) -> Path:
    source = _xlsx_file(tmp_path, 0.3)
    target = tmp_path / "float-artifact.xlsx"
    with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(target, "w") as target_zip:
        for item in source_zip.infolist():
            payload = source_zip.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b"<v>0.3</v>", b"<v>0.30000000000000004</v>")
            target_zip.writestr(item, payload)
    return target


def _xls_file(
    tmp_path: Path,
    value,
    *,
    number_format: str | None = None,
    name: str = "Typed XLS",
    email: str = "typed.xls@example.com",
) -> Path:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Alunos")
    for column, header in enumerate(("Aluno", "E-mail", "Matricula")):
        sheet.write(0, column, header)
    sheet.write(1, 0, name)
    sheet.write(1, 1, email)
    style = xlwt.easyxf(num_format_str=number_format) if number_format else xlwt.Style.default_style
    sheet.write(1, 2, value, style)
    path = tmp_path / f"typed-{uuid.uuid4().hex}.xls"
    workbook.save(str(path))
    return path


def test_m1_non_student_account_email_collision_rejects_without_mutation():
    _, _, turma_id = _seed_school()
    email = f"admin-{uuid.uuid4().hex}@example.com"
    user_id, _ = _insert_identity(name="Existing Admin", email=email, user_type="admin")

    _, flashes = _post_modal(turma_id, [("Imported Student", email, "M1.001")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT 1 FROM alunos WHERE matricula='M1.001'").fetchone() is None
        assert dict(conn.execute("SELECT nome,email,tipo FROM usuarios WHERE id=?", (user_id,)).fetchone()) == {
            "nome": "Existing Admin",
            "email": email,
            "tipo": "admin",
        }
    assert any(category == "error" and "conta" in message.lower() for category, message in flashes)


def test_m2_update_keeps_linked_usuario_and_aluno_identity_synchronized():
    _, _, turma_id = _seed_school()
    old_email = f"old-{uuid.uuid4().hex}@example.com"
    new_email = f"new-{uuid.uuid4().hex}@example.com"
    user_id, aluno_id = _insert_identity(
        name="Old Name",
        email=old_email,
        user_type="aluno",
        matricula="M2.001",
    )

    _post_modal(turma_id, [("New Name", new_email, "M2.001")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        user = conn.execute("SELECT nome,email,tipo FROM usuarios WHERE id=?", (user_id,)).fetchone()
        aluno = conn.execute(
            "SELECT usuario_id,nome,email,turma_id,status FROM alunos WHERE id=?", (aluno_id,)
        ).fetchone()
        assert dict(user) == {"nome": "New Name", "email": new_email, "tipo": "aluno"}
        assert dict(aluno) == {
            "usuario_id": user_id,
            "nome": "New Name",
            "email": new_email,
            "turma_id": turma_id,
            "status": "Ativo",
        }
        assert conn.execute("SELECT COUNT(*) FROM usuarios WHERE id=?", (user_id,)).fetchone()[0] == 1


def test_m2_conflicting_target_email_rejects_without_merging_students():
    _, _, turma_id = _seed_school()
    email_a = f"a-{uuid.uuid4().hex}@example.com"
    email_b = f"b-{uuid.uuid4().hex}@example.com"
    user_a, aluno_a = _insert_identity(
        name="Student A", email=email_a, user_type="aluno", matricula="M2.A"
    )
    user_b, aluno_b = _insert_identity(
        name="Student B", email=email_b, user_type="aluno", matricula="M2.B"
    )

    _, flashes = _post_modal(turma_id, [("Attempted Merge", email_b, "M2.A")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert dict(conn.execute("SELECT nome,email FROM usuarios WHERE id=?", (user_a,)).fetchone()) == {
            "nome": "Student A",
            "email": email_a,
        }
        assert dict(conn.execute("SELECT nome,email,usuario_id FROM alunos WHERE id=?", (aluno_a,)).fetchone()) == {
            "nome": "Student A",
            "email": email_a,
            "usuario_id": user_a,
        }
        assert conn.execute("SELECT usuario_id FROM alunos WHERE id=?", (aluno_b,)).fetchone()[0] == user_b
    assert any(category == "error" and "conflito" in message.lower() for category, message in flashes)


def test_identity_match_by_email_updates_same_student_without_duplicate_user():
    _, _, turma_id = _seed_school()
    email = f"email-match-{uuid.uuid4().hex}@example.com"
    user_id, aluno_id = _insert_identity(
        name="Email Match Old",
        email=email,
        user_type="aluno",
        matricula="EMAIL.OLD",
    )

    _post_modal(turma_id, [("Email Match New", email, "EMAIL.NEW")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert dict(conn.execute("SELECT nome,email FROM usuarios WHERE id=?", (user_id,)).fetchone()) == {
            "nome": "Email Match New",
            "email": email,
        }
        assert dict(
            conn.execute(
                "SELECT usuario_id,nome,email,matricula,turma_id,status FROM alunos WHERE id=?",
                (aluno_id,),
            ).fetchone()
        ) == {
            "usuario_id": user_id,
            "nome": "Email Match New",
            "email": email,
            "matricula": "EMAIL.NEW",
            "turma_id": turma_id,
            "status": "Ativo",
        }
        assert conn.execute("SELECT COUNT(*) FROM usuarios WHERE LOWER(email)=LOWER(?)", (email,)).fetchone()[0] == 1


@pytest.mark.parametrize(
    ("on_conflict", "expected_name", "expected_status"),
    [("skip", "Inactive Old", "Inativo"), ("update", "Inactive New", "Ativo")],
)
def test_existing_inactive_student_in_target_obeys_safe_conflict_strategy(
    on_conflict, expected_name, expected_status
):
    _, _, turma_id = _seed_school()
    email = f"inactive-{uuid.uuid4().hex}@example.com"
    matricula = f"INACTIVE.{uuid.uuid4().hex[:8]}"
    user_id, aluno_id = _insert_identity(
        name="Inactive Old",
        email=email,
        user_type="aluno",
        matricula=matricula,
        turma_id=turma_id,
        status="Inativo",
    )

    _post_modal(
        turma_id,
        [("Inactive New", email, matricula)],
        on_conflict=on_conflict,
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert dict(conn.execute("SELECT nome,email,tipo FROM usuarios WHERE id=?", (user_id,)).fetchone()) == {
            "nome": expected_name,
            "email": email,
            "tipo": "aluno",
        }
        assert dict(
            conn.execute(
                "SELECT usuario_id,nome,status,turma_id FROM alunos WHERE id=?", (aluno_id,)
            ).fetchone()
        ) == {
            "usuario_id": user_id,
            "nome": expected_name,
            "status": expected_status,
            "turma_id": turma_id,
        }


def test_orphan_student_account_is_rejected_without_silent_reconstruction():
    _, _, turma_id = _seed_school()
    email = f"orphan-user-{uuid.uuid4().hex}@example.com"
    user_id, _ = _insert_identity(name="Orphan User", email=email, user_type="aluno")

    _, flashes = _post_modal(turma_id, [("Imported", email, "ORPHAN.USER")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT 1 FROM alunos WHERE usuario_id=?", (user_id,)).fetchone() is None
    assert any(category == "error" and "correção manual" in message for category, message in flashes)


def test_student_without_linked_user_is_rejected_fail_closed():
    _, _, turma_id = _seed_school()
    email = f"orphan-student-{uuid.uuid4().hex}@example.com"
    with main.app.app_context():
        conn = app_db.get_db_connection()
        aluno_id = conn.execute(
            """
            INSERT INTO alunos (usuario_id,nome,email,matricula,status)
            VALUES (NULL,'Orphan Student',?,'ORPHAN.STUDENT','Ativo') RETURNING id
            """,
            (email,),
        ).fetchone()["id"]
        conn.commit()

    _, flashes = _post_modal(turma_id, [("Attempted Update", email, "ORPHAN.STUDENT")])

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert dict(
            conn.execute("SELECT usuario_id,nome,turma_id FROM alunos WHERE id=?", (aluno_id,)).fetchone()
        ) == {"usuario_id": None, "nome": "Orphan Student", "turma_id": None}
    assert any(category == "error" and "não possui conta" in message for category, message in flashes)


@pytest.mark.parametrize(
    ("value", "expected", "number_format"),
    [(3.206, "3.206", "General"), (3.26, "3.26", "General"), (123, "123", "General"), (123, "00123", "00000")],
)
def test_m3_xlsx_safe_numeric_identifiers(value, expected, number_format, tmp_path):
    row = parse_student_import(_xlsx_file(tmp_path, value, number_format=number_format))[0]
    assert row.matricula == expected


@pytest.mark.parametrize("matricula", ["3.206", "3.26", "00123"])
def test_m3_csv_preserves_textual_identifiers(matricula, tmp_path):
    path = tmp_path / "safe.csv"
    path.write_text(
        f"Aluno,E-mail,Matricula\nText,text@example.com,{matricula}\n",
        encoding="utf-8",
    )
    assert parse_student_import(path)[0].matricula == matricula


@pytest.mark.parametrize(
    ("value", "data_type", "message"),
    [
        (dt.datetime(2026, 1, 2, 3, 4, 5), None, "data"),
        ("#DIV/0!", "e", "erro"),
        ("=1+1", "f", "fórmula"),
    ],
)
def test_m3_xlsx_unsafe_identifier_types_are_rejected(value, data_type, message, tmp_path):
    path = _xlsx_file(tmp_path, value, data_type=data_type)
    with pytest.raises(StudentImportError, match=message):
        parse_student_import(path)


def test_m3_xlsx_rejects_unsafe_float_artifact(tmp_path):
    with pytest.raises(StudentImportError, match="numérico"):
        parse_student_import(_xlsx_float_artifact_file(tmp_path))


@pytest.mark.parametrize(
    "value",
    [float("nan"), Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity")],
)
def test_m3_rejects_non_finite_numeric_identifier(value):
    with pytest.raises(StudentImportError, match="numérico inválido"):
        normalize_student_import_values(
            [("Nonfinite", "nonfinite@example.com", value)]
        )


def test_m3_xls_rejects_date_cell_and_preserves_explicit_zero_padding(tmp_path):
    date_path = _xls_file(tmp_path, dt.datetime(2026, 1, 2), number_format="YYYY-MM-DD")
    padded_path = _xls_file(tmp_path, 123, number_format="00000")

    with pytest.raises(StudentImportError, match="data"):
        parse_student_import(date_path)
    assert parse_student_import(padded_path)[0].matricula == "00123"


@pytest.mark.parametrize(
    ("value", "expected"),
    [(3.206, "3.206"), (3.26, "3.26"), (123, "123")],
)
def test_m3_xls_safe_numeric_identifiers(value, expected, tmp_path):
    assert parse_student_import(_xls_file(tmp_path, value))[0].matricula == expected


@pytest.mark.parametrize("mode", ["add", "edit"])
@pytest.mark.parametrize("duplicate", ["matricula", "email"])
def test_m4_live_add_edit_reject_import_duplicates(mode, duplicate):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode == "edit"))
    email_a = f"dup-{uuid.uuid4().hex}@example.com"
    rows = [
        ("First", email_a, "M4.DUP", "ATIVO"),
        (
            "Second",
            email_a.upper() if duplicate == "email" else f"other-{uuid.uuid4().hex}@example.com",
            "M4.OTHER" if duplicate == "email" else "M4.DUP",
            "ATIVO",
        ),
    ]

    response, flashes = _post_turma_form(mode, curso_id, matriz_id, turma_id, rows)

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT 1 FROM alunos WHERE matricula IN ('M4.DUP','M4.OTHER')").fetchone() is None
        if mode == "add":
            assert conn.execute("SELECT COUNT(*) FROM turmas WHERE curso_id=?", (curso_id,)).fetchone()[0] == 0
    feedback = str(response).lower() + " ".join(
        message.lower() for _, message in flashes
    )
    assert "duplic" in feedback


@pytest.mark.parametrize("duplicate", ["matricula", "email"])
def test_m4_modal_import_rejects_duplicates(duplicate):
    _, _, turma_id = _seed_school()
    email = f"modal-dup-{uuid.uuid4().hex}@example.com"
    rows = [
        ("First", email, "MODAL.DUP"),
        (
            "Second",
            email.upper() if duplicate == "email" else f"other-{uuid.uuid4().hex}@example.com",
            "MODAL.OTHER" if duplicate == "email" else "MODAL.DUP",
        ),
    ]

    _, flashes = _post_modal(turma_id, rows)

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute(
            "SELECT 1 FROM alunos WHERE matricula IN ('MODAL.DUP','MODAL.OTHER')"
        ).fetchone() is None
    assert any(category == "error" and "duplic" in message.lower() for category, message in flashes)


@pytest.mark.parametrize("mode", ["add", "edit"])
def test_m4_live_add_edit_reject_split_student_identity(mode):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode == "edit"))
    email_a = f"split-a-{uuid.uuid4().hex}@example.com"
    email_b = f"split-b-{uuid.uuid4().hex}@example.com"
    matricula_a = f"SPLIT.A.{uuid.uuid4().hex[:8]}"
    matricula_b = f"SPLIT.B.{uuid.uuid4().hex[:8]}"
    user_a, aluno_a = _insert_identity(
        name="Split A", email=email_a, user_type="aluno", matricula=matricula_a
    )
    user_b, aluno_b = _insert_identity(
        name="Split B", email=email_b, user_type="aluno", matricula=matricula_b
    )

    response, flashes = _post_turma_form(
        mode,
        curso_id,
        matriz_id,
        turma_id,
        [("Invalid Merge", email_b, matricula_a, "ATIVO")],
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT usuario_id FROM alunos WHERE id=?", (aluno_a,)).fetchone()[0] == user_a
        assert conn.execute("SELECT usuario_id FROM alunos WHERE id=?", (aluno_b,)).fetchone()[0] == user_b
        if mode == "add":
            assert conn.execute("SELECT COUNT(*) FROM turmas WHERE curso_id=?", (curso_id,)).fetchone()[0] == 0
    feedback = str(response).lower() + " ".join(message.lower() for _, message in flashes)
    assert "conflito" in feedback


@pytest.mark.parametrize("mode", ["add", "edit"])
def test_m4_live_add_edit_import_forces_active_status(mode):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode == "edit"))
    matricula = f"M4.ACTIVE.{uuid.uuid4().hex[:8]}"
    email = f"active-{uuid.uuid4().hex}@example.com"

    _, flashes = _post_turma_form(
        mode, curso_id, matriz_id, turma_id, [("Active", email, matricula, "INATIVO")]
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        row = conn.execute("SELECT status FROM alunos WHERE matricula=?", (matricula,)).fetchone()
        assert row is not None, flashes
        assert row[0] == "Ativo"


@pytest.mark.parametrize("mode", ["add", "edit"])
def test_m4_live_add_edit_server_rejects_date_typed_xlsx_matricula(mode, tmp_path):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode == "edit"))
    spreadsheet = _xlsx_file(tmp_path, dt.datetime(2026, 1, 2))
    preview_matricula = f"PREVIEW.{uuid.uuid4().hex[:8]}"

    response, flashes = _post_turma_form(
        mode,
        curso_id,
        matriz_id,
        turma_id,
        [("Preview", "preview@example.com", preview_matricula, "ATIVO")],
        import_file=(spreadsheet.read_bytes(), "date-typed.xlsx"),
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT 1 FROM alunos WHERE matricula=?", (preview_matricula,)).fetchone() is None
        assert conn.execute("SELECT 1 FROM alunos WHERE email='typed.xlsx@example.com'").fetchone() is None
        if mode == "add":
            assert conn.execute("SELECT COUNT(*) FROM turmas WHERE curso_id=?", (curso_id,)).fetchone()[0] == 0
    feedback = str(response).lower() + " ".join(message.lower() for _, message in flashes)
    assert "data" in feedback


@pytest.mark.parametrize("mode", ["add", "edit"])
@pytest.mark.parametrize("extension", ["xlsx", "xls"])
def test_m4_live_add_edit_server_imports_real_spreadsheet_file(mode, extension, tmp_path):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode == "edit"))
    matricula = f"LIVE.{extension.upper()}.{uuid.uuid4().hex[:8]}"
    email = f"live-{extension}-{uuid.uuid4().hex}@example.com"
    name = f"Live {extension.upper()}"
    spreadsheet = (
        _xlsx_file(tmp_path, matricula, name=name, email=email)
        if extension == "xlsx"
        else _xls_file(tmp_path, matricula, name=name, email=email)
    )

    _post_turma_form(
        mode,
        curso_id,
        matriz_id,
        turma_id,
        [("Preview", "preview@example.com", "PREVIEW.IGNORED", "INATIVO")],
        import_file=(spreadsheet.read_bytes(), f"students.{extension}"),
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        row = conn.execute(
            """
            SELECT a.nome,a.email,a.status,a.turma_id,u.nome AS user_name,u.email AS user_email,u.tipo
              FROM alunos a JOIN usuarios u ON u.id=a.usuario_id
             WHERE a.matricula=?
            """,
            (matricula,),
        ).fetchone()
        assert dict(row) == {
            "nome": name,
            "email": email,
            "status": "Ativo",
            "turma_id": turma_id if mode == "edit" else row["turma_id"],
            "user_name": name,
            "user_email": email,
            "tipo": "aluno",
        }
        assert row["turma_id"] is not None


def test_m4_edit_import_updates_student_already_rendered_in_roster():
    curso_id, matriz_id, turma_id = _seed_school(with_turma=True)
    email = f"existing-roster-{uuid.uuid4().hex}@example.com"
    matricula = f"ROSTER.{uuid.uuid4().hex[:8]}"
    user_id, aluno_id = _insert_identity(
        name="Roster Old",
        email=email,
        user_type="aluno",
        matricula=matricula,
        turma_id=turma_id,
        status="Inativo",
    )
    body = (
        "Aluno,E-mail,Matricula\n"
        f"Roster Updated,{email},{matricula}\n"
    ).encode("utf-8")

    _, flashes = _post_turma_form(
        "edit",
        curso_id,
        matriz_id,
        turma_id,
        [
            ("Roster Old", email, matricula, "INATIVO"),
            ("Roster Updated", email, matricula, "ATIVO"),
        ],
        import_file=(body, "existing-roster.csv"),
        imported_flags=["0", "1"],
    )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert dict(
            conn.execute(
                """
                SELECT a.id,a.nome,a.email,a.matricula,a.turma_id,a.status,
                       u.id AS user_id,u.nome AS user_name,u.email AS user_email
                  FROM alunos a JOIN usuarios u ON u.id=a.usuario_id
                 WHERE a.id=?
                """,
                (aluno_id,),
            ).fetchone()
        ) == {
            "id": aluno_id,
            "nome": "Roster Updated",
            "email": email,
            "matricula": matricula,
            "turma_id": turma_id,
            "status": "Ativo",
            "user_id": user_id,
            "user_name": "Roster Updated",
            "user_email": email,
        }
        assert conn.execute(
            "SELECT COUNT(*) FROM alunos WHERE matricula=?", (matricula,)
        ).fetchone()[0] == 1
    assert not any(category == "error" for category, _ in flashes), flashes


@pytest.mark.parametrize("mode", ["modal", "add", "edit"])
def test_true_later_row_identity_failure_rolls_back_earlier_mutation(mode):
    curso_id, matriz_id, turma_id = _seed_school(with_turma=(mode != "add"))
    admin_email = f"rollback-admin-{uuid.uuid4().hex}@example.com"
    admin_id, _ = _insert_identity(name="Rollback Admin", email=admin_email, user_type="admin")
    first_matricula = f"ROLLBACK.{uuid.uuid4().hex[:8]}"
    first_email = f"rollback-first-{uuid.uuid4().hex}@example.com"
    rows = [
        ("First Write", first_email, first_matricula),
        ("Forbidden Merge", admin_email, f"ROLLBACK.ADMIN.{uuid.uuid4().hex[:8]}"),
    ]

    if mode == "modal":
        _, flashes = _post_modal(int(turma_id), rows)
    else:
        response, flashes = _post_turma_form(
            mode,
            curso_id,
            matriz_id,
            turma_id,
            [(name, email, matricula, "ATIVO") for name, email, matricula in rows],
        )

    with main.app.app_context():
        conn = app_db.get_db_connection()
        assert conn.execute("SELECT 1 FROM alunos WHERE matricula=?", (first_matricula,)).fetchone() is None
        assert dict(conn.execute("SELECT nome,email,tipo FROM usuarios WHERE id=?", (admin_id,)).fetchone()) == {
            "nome": "Rollback Admin",
            "email": admin_email,
            "tipo": "admin",
        }
        if mode == "add":
            assert conn.execute("SELECT COUNT(*) FROM turmas WHERE curso_id=?", (curso_id,)).fetchone()[0] == 0
    if mode == "modal":
        assert any(category == "error" for category, _ in flashes), flashes
    else:
        assert "erro" in str(response).lower()
