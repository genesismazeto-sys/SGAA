from __future__ import annotations

import io
import uuid
from pathlib import Path

import openpyxl
import pytest
import xlwt
from flask import session

from app import db as app_db
from app.views.admin import alunos_turmas_cursos as target
import main


def _seed_turma() -> int:
    suffix = uuid.uuid4().hex[:10]
    with main.app.app_context():
        conn = app_db.get_db_connection()
        curso_id = conn.execute(
            "INSERT INTO cursos (nome,codigo,duracao_periodos,status) VALUES (?,?,8,'ativo') RETURNING id",
            (f"Curso import {suffix}", f"IMP-{suffix}"),
        ).fetchone()["id"]
        turma_id = conn.execute(
            """
            INSERT INTO turmas (nome,codigo,numero,curso_id,status,ano_inicio,semestre_inicio)
            VALUES (?,?,?,?, 'Ativa', 2026, 1) RETURNING id
            """,
            (f"Turma import {suffix}", f"IMP-{suffix}-T1", 1, curso_id),
        ).fetchone()["id"]
        conn.commit()
        return int(turma_id)


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _xls_bytes(rows: list[list[object]]) -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Alunos")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            sheet.write(row_index, column_index, value)
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _post_import(turma_id: int, payload: bytes, filename: str):
    with main.app.test_request_context(
        "/admin/turmas/importar",
        method="POST",
        data={
            "turma_id": str(turma_id),
            "csv_arquivo": (io.BytesIO(payload), filename),
        },
    ):
        session["user_id"] = 999_999
        session["user_type"] = "admin"
        response = target.admin_turmas_importar()
        flashes = list(session.get("_flashes", []))
        return response, flashes


def _student_by_matricula(matricula: str):
    with main.app.app_context():
        conn = app_db.get_db_connection()
        return conn.execute(
            "SELECT nome,email,matricula,turma_id,status FROM alunos WHERE matricula=?",
            (matricula,),
        ).fetchone()


def _delete_students(*matriculas: str) -> None:
    with main.app.app_context():
        conn = app_db.get_db_connection()
        placeholders = ",".join("?" for _ in matriculas)
        user_ids = [
            row["usuario_id"]
            for row in conn.execute(
                f"SELECT usuario_id FROM alunos WHERE matricula IN ({placeholders})",
                matriculas,
            ).fetchall()
            if row["usuario_id"] is not None
        ]
        conn.execute(f"DELETE FROM alunos WHERE matricula IN ({placeholders})", matriculas)
        if user_ids:
            user_placeholders = ",".join("?" for _ in user_ids)
            conn.execute(f"DELETE FROM usuarios WHERE id IN ({user_placeholders})", user_ids)
        conn.commit()


def test_csv_canonical_order_creates_active_student_and_guards_positions():
    turma_id = _seed_turma()
    payload = (
        "Aluno,E-mail,Matricula\n"
        "Alexandre Posenatto,alexandre.posenatto@example.com,3.206\n"
    ).encode("utf-8")

    try:
        response, _ = _post_import(turma_id, payload, "canonical.csv")

        assert response.status_code == 302
        row = _student_by_matricula("3.206")
        assert row is not None
        assert dict(row) == {
            "nome": "Alexandre Posenatto",
            "email": "alexandre.posenatto@example.com",
            "matricula": "3.206",
            "turma_id": turma_id,
            "status": "Ativo",
        }
    finally:
        _delete_students("3.206")


def test_xlsx_uses_same_three_column_contract_and_preserves_matricula_text():
    turma_id = _seed_turma()
    payload = _xlsx_bytes(
        [
            ["Aluno", "E-mail", "Matricula"],
            ["Numeric Dot A", "numeric.a@example.com", 3.206],
            ["Numeric Dot B", "numeric.b@example.com", 3.26],
            ["Leading Zero", "leading.zero@example.com", "00123"],
        ]
    )

    response, _ = _post_import(turma_id, payload, "canonical.xlsx")

    assert response.status_code == 302
    assert _student_by_matricula("3.206")["nome"] == "Numeric Dot A"
    assert _student_by_matricula("3.26")["nome"] == "Numeric Dot B"
    assert _student_by_matricula("00123")["nome"] == "Leading Zero"
    assert _student_by_matricula("00123")["status"] == "Ativo"


def test_real_legacy_xls_imports_with_active_status():
    turma_id = _seed_turma()
    payload = _xls_bytes(
        [
            ["Aluno", "E-mail", "Matricula"],
            ["Legacy XLS Student", "legacy.xls.student@example.com", "XLS.003"],
        ]
    )

    response, _ = _post_import(turma_id, payload, "canonical.xls")

    assert response.status_code == 302
    row = _student_by_matricula("XLS.003")
    assert row is not None
    assert row["nome"] == "Legacy XLS Student"
    assert row["email"] == "legacy.xls.student@example.com"
    assert row["status"] == "Ativo"


def test_obsolete_order_is_rejected_without_positional_corruption():
    turma_id = _seed_turma()
    payload = (
        "Matricula,Aluno,E-mail\n"
        "OLD.001,Obsolete Order,obsolete.order@example.com\n"
    ).encode("utf-8")

    response, flashes = _post_import(turma_id, payload, "obsolete.csv")

    assert response.status_code == 302
    assert _student_by_matricula("OLD.001") is None
    assert any(category == "error" and "Aluno, E-mail, Matricula" in message for category, message in flashes)


@pytest.mark.parametrize("filename", ["malformed.xlsx", "malformed.xls"])
def test_malformed_spreadsheet_returns_controlled_feedback_and_writes_nothing(filename):
    turma_id = _seed_turma()
    marker = f"MAL-{uuid.uuid4().hex[:10]}"

    response, flashes = _post_import(turma_id, b"not-a-spreadsheet", filename)

    assert response.status_code == 302
    assert _student_by_matricula(marker) is None
    assert any(category == "error" and "planilha" in message.lower() for category, message in flashes)


def test_header_only_file_is_rejected_and_status_column_is_not_required():
    turma_id = _seed_turma()

    response, flashes = _post_import(
        turma_id,
        b"Aluno,E-mail,Matricula\n",
        "header-only.csv",
    )

    assert response.status_code == 302
    assert any(category == "error" and "nenhum aluno" in message.lower() for category, message in flashes)


def test_row_validation_is_all_or_nothing():
    turma_id = _seed_turma()
    marker = f"ROLLBACK-{uuid.uuid4().hex[:10]}"
    payload = (
        "Aluno,E-mail,Matricula\n"
        f"Valid Before Error,valid.before.error@example.com,{marker}\n"
        "Missing Email,,MISSING.002\n"
    ).encode("utf-8")

    response, flashes = _post_import(turma_id, payload, "atomic.csv")

    assert response.status_code == 302
    assert _student_by_matricula(marker) is None
    assert any(category == "error" and "Linha 3" in message for category, message in flashes)


def test_extra_column_and_unsupported_extension_are_rejected():
    turma_id = _seed_turma()
    extra_payload = (
        "Aluno,E-mail,Matricula,Situação\n"
        "Fourth Field,fourth.field@example.com,FOURTH.001,Ativo\n"
    ).encode("utf-8")

    _, extra_flashes = _post_import(turma_id, extra_payload, "four-columns.csv")
    _, unsupported_flashes = _post_import(turma_id, b"payload", "students.ods")

    assert _student_by_matricula("FOURTH.001") is None
    assert any(category == "error" and "exatamente" in message for category, message in extra_flashes)
    assert any(category == "error" and "CSV, XLSX ou XLS" in message for category, message in unsupported_flashes)


def test_ui_advertises_only_the_canonical_three_column_contract():
    project_root = Path(__file__).resolve().parents[1]
    paths = [
        project_root / "templates" / "admin_turmas.html",
        project_root / "templates" / "admin_turma_alunos.html",
        project_root / "templates" / "admin_turma_form.html",
        project_root / "templates" / "admin_importar_turma.html",
        project_root / "templates" / "admin_adicionar_turma.html",
        project_root / "templates" / "admin_editar_turma.html",
        project_root / "templates" / "partials" / "admin_importar_turma_form.html",
    ]
    combined = "\n".join(path.read_text(encoding="utf-8") for path in paths)
    preview_source = (
        project_root / "static" / "js" / "student-import-preview.js"
    ).read_text(encoding="utf-8")

    assert "Importar Alunos (CSV/XLSX/XLS)" in combined
    assert "Aluno, E-mail, Matricula" in combined
    assert 'accept=".csv,.xlsx,.xls"' in combined
    assert "matricula,nome,email" not in combined.lower()
    assert combined.count("StudentImportPreview.read") == 2
    assert combined.count('name="student_import_file"') == 2
    assert combined.count('enctype="multipart/form-data"') >= 2
    assert combined.count('name="aluno_importado[]"') >= 3
    assert "fileImp.value = ''" not in combined
    assert "parseCSVFile" not in combined
    assert "sheet_to_json" not in combined
    assert "sheet_to_json" in preview_source


def test_live_add_and_edit_pages_own_local_reader_and_nonblocking_fallback():
    turma_id = _seed_turma()
    with main.app.app_context():
        conn = app_db.get_db_connection()
        admin = conn.execute(
            "SELECT id, nome FROM usuarios WHERE tipo='admin' ORDER BY id LIMIT 1"
        ).fetchone()
        assert admin is not None

    client = main.app.test_client()
    with client.session_transaction() as flask_session:
        flask_session["user_id"] = int(admin["id"])
        flask_session["user_name"] = admin["nome"]
        flask_session["user_type"] = "admin"

    reader_response = client.get("/static/vendor/xlsx.full.min.js")
    assert reader_response.status_code == 200

    responses = (
        client.get("/admin/adicionar_turma"),
        client.get(f"/admin/editar_turma/{turma_id}"),
    )
    for response in responses:
        assert response.status_code == 200
        html = response.get_data(as_text=True)
        local_reader = "/static/vendor/xlsx.full.min.js"
        preview = "/static/js/student-import-preview.js"
        assert local_reader in html
        assert "cdn.jsdelivr.net/npm/xlsx" not in html
        assert html.index(local_reader) < html.index(preview)
        assert "StudentImportPreview.isAvailable()" in html
        assert "Pré-visualização indisponível; o arquivo será processado ao salvar." in html

    modal_response = client.get("/admin/turmas")
    assert modal_response.status_code == 200
    modal_html = modal_response.get_data(as_text=True)
    assert f'<option value="{turma_id}">' in modal_html
